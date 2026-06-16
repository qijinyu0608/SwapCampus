#!/usr/bin/env python3

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "docs" / "templates"
OUTPUT_DIR = ROOT / "docs" / "generated-docx"
IMAGE_DIR = ROOT / "artifacts" / "screenshots"

BODY_FONT_CN = "宋体"
BODY_FONT_EN = "Times New Roman"
BODY_SIZE = Pt(12)
SMALL_SIZE = Pt(10.5)
DATE_TEXT = "2026-06-15"


def set_run_font(run, *, size=BODY_SIZE, bold=False, italic=False) -> None:
    run.font.name = BODY_FONT_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_CN)
    run.font.size = size
    run.bold = bold
    run.italic = italic


def style_paragraph(
    paragraph,
    *,
    size=BODY_SIZE,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_indent=True,
    left_indent_cm=0.0,
) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.35
    fmt.left_indent = Cm(left_indent_cm) if left_indent_cm else None
    fmt.first_line_indent = Cm(0.74) if first_line_indent else None
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def clear_cell(cell) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.text = ""


def write_paragraph(
    paragraph,
    text: str,
    *,
    size=BODY_SIZE,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_indent=True,
    left_indent_cm=0.0,
) -> None:
    paragraph.text = text
    style_paragraph(
        paragraph,
        size=size,
        bold=bold,
        align=align,
        first_line_indent=first_line_indent,
        left_indent_cm=left_indent_cm,
    )


def fill_text_cell(cell, blocks: list[dict | str]) -> None:
    clear_cell(cell)
    first = True
    for block in blocks:
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        if isinstance(block, str):
            write_paragraph(paragraph, block)
            continue

        kind = block.get("type", "p")
        if kind == "p":
            write_paragraph(
                paragraph,
                block["text"],
                size=block.get("size", BODY_SIZE),
                bold=block.get("bold", False),
                align=block.get("align", WD_ALIGN_PARAGRAPH.LEFT),
                first_line_indent=block.get("first_line_indent", True),
                left_indent_cm=block.get("left_indent_cm", 0.0),
            )
        elif kind == "bullet":
            write_paragraph(
                paragraph,
                f"- {block['text']}",
                size=block.get("size", BODY_SIZE),
                first_line_indent=False,
                left_indent_cm=0.74,
            )
        elif kind == "number":
            write_paragraph(
                paragraph,
                f"{block['index']}. {block['text']}",
                size=block.get("size", BODY_SIZE),
                first_line_indent=False,
                left_indent_cm=0.74,
            )
        elif kind == "image":
            paragraph.alignment = block.get("align", WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run()
            run.add_picture(str(block["path"]), width=block.get("width", Cm(13.5)))
            if block.get("caption"):
                caption = cell.add_paragraph()
                write_paragraph(
                    caption,
                    block["caption"],
                    size=SMALL_SIZE,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    first_line_indent=False,
                )
        elif kind == "blank":
            write_paragraph(
                paragraph,
                "",
                size=block.get("size", BODY_SIZE),
                first_line_indent=False,
            )
        else:
            raise ValueError(f"Unsupported block type: {kind}")


def set_table_row(table, row_index: int, values: list[str]) -> None:
    row = table.rows[row_index]
    for idx, value in enumerate(values):
        cell = row.cells[idx]
        clear_cell(cell)
        write_paragraph(
            cell.paragraphs[0],
            value,
            first_line_indent=False,
            size=BODY_SIZE,
        )


def fill_cover(table, *, code: str, name: str) -> None:
    values = {
        "编号": code,
        "文件名": name,
        "类型": "产出",
        "版本": "v1.0 · 2026-06-15",
        "小组 / 组号": "待补充",
        "班级 / 学号 / 姓名": "待补充 / 待补充 / 待补充",
        "完成日期": DATE_TEXT,
    }
    for row in table.rows:
        key = row.cells[0].text.strip()
        if key in values:
            clear_cell(row.cells[1])
            write_paragraph(
                row.cells[1].paragraphs[0],
                values[key],
                first_line_indent=False,
            )


def set_heading(paragraph, text: str) -> None:
    paragraph.text = text
    style_paragraph(paragraph, size=BODY_SIZE, bold=False, first_line_indent=False)


def write_readme(diagram_items: list[str]) -> None:
    text = [
        "# 课程设计 DOCX 生成结果",
        "",
        "本目录保存基于 `docs/templates/` 模板和当前 Markdown 内容生成的正式 `docx` 文档。",
        "",
        "## 本地生成",
        "",
        "先安装 Python 依赖：",
        "",
        "```bash",
        "python3 -m pip install --target .docs-py -r requirements-docs.txt",
        "```",
        "",
        "再执行生成命令：",
        "",
        "```bash",
        "make docs-docx",
        "```",
        "",
        "也可以直接运行：",
        "",
        "```bash",
        "PYTHONPATH=.docs-py python3 scripts/generate_course_docx.py",
        "```",
        "",
        "生成结果会覆盖当前目录下的 `D-01` 到 `D-09` 文档，并刷新本说明文件。",
        "",
        "## GitHub Actions",
        "",
        "仓库中的 `.github/workflows/ci.yml` 会在 `push`、`pull_request` 和手动 `workflow_dispatch` 时执行：",
        "",
        "1. 后端测试",
        "2. 前端构建",
        "3. 课程设计 `docx` 文档生成",
        "",
        "文档生成 job 会上传名为 `course-docx` 的 artifact，内容即本目录下的生成结果。",
        "",
        "## 文件清单",
        "",
        "- `D-01_项目开题报告.docx`",
        "- `D-02_需求规格说明书_SRS.docx`",
        "- `D-03_概要设计说明书.docx`",
        "- `D-04_详细设计说明书.docx`",
        "- `D-05_数据库设计说明书.docx`",
        "- `D-06_软件测试计划与测试报告.docx`",
        "- `D-07_用户手册.docx`",
        "- `D-08_部署与运维手册.docx`",
        "- `D-09_课程设计总结报告.docx`",
        "",
        "## 待补图示清单",
        "",
    ]
    for item in diagram_items:
        text.append(f"- {item}")
    (OUTPUT_DIR / "README.md").write_text("\n".join(text) + "\n", encoding="utf-8")


def build_d01() -> None:
    doc = Document(TEMPLATE_DIR / "D-01_项目开题报告_模板.docx")
    fill_cover(doc.tables[0], code="D-01", name="项目开题报告")

    basic = doc.tables[2]
    mapping = {
        "项目中文名": "SwapCampus 校园闲置物品交易平台",
        "项目英文名": "SwapCampus Campus Idle Goods Trading Platform",
        "选题编号": "T-02",
        "小组组号": "待补充",
        "指导教师": "待补充",
        "计划起讫": "2026-06-02 至 2026-06-15",
    }
    for row in basic.rows[1:]:
        label = row.cells[0].text.strip()
        if label in mapping:
            clear_cell(row.cells[1])
            write_paragraph(row.cells[1].paragraphs[0], mapping[label], first_line_indent=False)

    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "本项目面向校内师生，聚焦教材、宿舍用品、数码配件、运动器材等校园高频闲置交易场景。项目采用前后端分离的全栈 TypeScript 技术路线，通过 Docker Compose 提供可复现部署环境，并以课程设计要求的文档、测试、验收和交付材料作为完整成果的一部分。",
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            "校内闲置物品交易长期依赖微信群、班群和熟人转发，信息入口分散、交易留痕不足、信用约束弱、售后责任模糊。SwapCampus 的目标是在校园这一明确边界内，提供一个可实名、可追溯、可审核、可复现部署的闲置交易平台，用较低的学习成本替代碎片化沟通方式，同时体现课程设计对需求分析、架构设计、数据库建模、测试与运维的完整软件工程训练价值。",
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "p", "text": "同类产品对比以“社会化二手平台、校园论坛/群组、垂直校园服务页”三类对象为主。"},
            {"type": "bullet", "text": "闲鱼：商品流和交易机制成熟，但主要面向社会化用户，无法天然提供校园身份认证、近场履约和校内治理规则。"},
            {"type": "bullet", "text": "微信群/QQ群：沟通成本低，但信息沉底快、检索弱、收藏与评价留痕缺失，无法支撑完整的订单闭环。"},
            {"type": "bullet", "text": "校园论坛/表白墙：具备一定校内传播能力，但通常只承载信息展示，不具备商品发布规则、订单、申诉、后台治理等工程化能力。"},
            {"type": "p", "text": "综合来看，现有方案要么偏泛化、要么偏临时沟通，尚未在“校内真实身份 + 校园高频商品结构 + 消息会话 + 交易留痕 + 后台治理”这一组合上形成完整闭环。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            {"type": "bullet", "text": "服务范围清晰：仅面向校内师生，功能边界明确，便于围绕近场交易优化产品路径。"},
            {"type": "bullet", "text": "信任机制更强：学号/邮箱注册、实名补充信息、信用摘要、评价和举报流程共同降低交易风险。"},
            {"type": "bullet", "text": "治理链路完整：支持商品审核、举报处理、实名审核、封禁/解封与审计日志，优于群聊式交易。"},
            {"type": "bullet", "text": "工程价值突出：项目涵盖对象存储、搜索、认证、消息推送、数据库建模、CI 与可复现部署。"},
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            "当前交付版已完成商品流、发布、订单、评价、申诉、消息、校园服务、信用中心和后台治理主链路，并提供 Docker 化启动与课程设计文档体系。完整 IM、更多前端自动化测试和更细粒度风控仍作为后续增强项。",
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            "一句话目标：构建一个面向校内师生、具备实名约束、交易留痕、消息沟通和后台治理能力的校园闲置交易平台。",
            {"type": "bullet", "text": "实现商品浏览、搜索、发布、收藏、下单、评价闭环。"},
            {"type": "bullet", "text": "建立基于学号实名、信用分和评价留痕的可信体系。"},
            {"type": "bullet", "text": "提供消息会话、校园服务和后台审核等扩展能力。"},
            {"type": "bullet", "text": "形成完整的设计、测试、验收和部署交付包。"},
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "Must：账号注册登录、商品发布与浏览、商品详情、订单创建、订单完成、评价、举报、后台商品审核。"},
            {"type": "bullet", "text": "Must：Docker Compose 启动、数据库建模、文档与代码同仓维护。"},
            {"type": "bullet", "text": "Should：消息会话、校园服务、收藏持久化、信用中心、实名审核、用户封禁。"},
            {"type": "bullet", "text": "Could：更细粒度推荐排序、更多运营看板、更多异常流程自动化验证。"},
            {"type": "bullet", "text": "Won't：本轮不接入真实支付、物流、跨校扩展与完整生产级 IM。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            {"type": "bullet", "text": "性能：本机演示环境下核心健康检查与商品读取接口保持秒级内响应，目标 P95 不高于 350ms。"},
            {"type": "bullet", "text": "可用性：通过 `make init` 与 `make start` 可在 1 小时内完成独立部署。"},
            {"type": "bullet", "text": "安全性：对注册、登录、会话、封禁拦截、媒体上传和后台治理做基础校验。"},
            {"type": "bullet", "text": "可维护性：采用模块化 NestJS、Prisma schema、Docker Compose 和 CI 基础流程。"},
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            "上述目标与范围遵循“先核心交易闭环、后扩展治理与体验”的迭代思路，确保课程设计阶段先交付可运行系统，再逐步补强体验、实时性和风控细节。",
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            {"type": "bullet", "text": "前端：React 18 + Vite，负责用户端、管理端和可演示页面交互。"},
            {"type": "bullet", "text": "后端：NestJS 单体模块化架构，便于在课程规模内保持清晰分层。"},
            {"type": "bullet", "text": "数据库：MySQL 8 承载业务主数据；SuperTokens 独立 PostgreSQL 承载认证会话。"},
            {"type": "bullet", "text": "中间件：MinIO 负责媒体存储，Meilisearch 负责搜索，Socket.IO 负责演示级实时推送。"},
            {"type": "bullet", "text": "部署：Docker Compose 统一管理容器依赖，降低演示环境复现成本。"},
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            "当前文档先以文字化里程碑替代甘特图截图，正式提交时建议补充一张按 D1-D14 展开的甘特图或项目计划截图。",
            {"type": "bullet", "text": "D1-D2：选题、开题、技术路线与仓库初始化。"},
            {"type": "bullet", "text": "D3-D6：注册登录、商品流、发布与订单主链路。"},
            {"type": "bullet", "text": "D7-D10：消息、校园服务、后台治理与对象存储补齐。"},
            {"type": "bullet", "text": "D11-D14：测试、验收、答辩材料和总结文档收口。"},
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            {"type": "bullet", "text": "里程碑 M1：基础工程可启动。"},
            {"type": "bullet", "text": "里程碑 M2：商品发布与订单主链路打通。"},
            {"type": "bullet", "text": "里程碑 M3：治理与扩展模块完成。"},
            {"type": "bullet", "text": "里程碑 M4：部署、测试、文档和答辩材料完成。"},
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            {"type": "bullet", "text": "组长：项目管理、进度控制、需求跟踪、答辩协调。"},
            {"type": "bullet", "text": "前端负责人：用户端与后台页面开发、联调、可视化表达。"},
            {"type": "bullet", "text": "后端负责人：接口、数据库、消息、治理逻辑与部署脚本。"},
            {"type": "bullet", "text": "测试文档负责人：测试执行、缺陷跟踪、文档同步、验收证据整理。"},
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            "当前项目采用“小步迭代、边实现边验收、边同步文档”的协作方式，避免开发尾声集中补文档带来的内容失真和交付风险。",
        ],
    )

    risk_table = doc.tables[17]
    set_table_row(risk_table, 1, ["R-01", "Docker 与多容器环境启动不稳定", "中", "高", "统一使用 Makefile 与 Compose 脚本，保留健康检查与重置命令", "待补充"])
    set_table_row(risk_table, 2, ["R-02", "课程周期内功能过多导致范围膨胀", "高", "高", "优先保证商品交易闭环，再逐步增加校园服务与治理能力", "待补充"])
    new_row = risk_table.add_row()
    for idx, value in enumerate(["R-03", "消息实时性与会话治理复杂度超出预期", "中", "中", "本轮先完成演示级推送，完整 IM 作为增强项", "待补充"]):
        clear_cell(new_row.cells[idx])
        write_paragraph(new_row.cells[idx].paragraphs[0], value, first_line_indent=False)
    new_row = risk_table.add_row()
    for idx, value in enumerate(["R-04", "图示与答辩材料准备滞后", "中", "中", "在文档生成后单独补齐架构图、ER 图、时序图与演示脚本", "待补充"]):
        clear_cell(new_row.cells[idx])
        write_paragraph(new_row.cells[idx].paragraphs[0], value, first_line_indent=False)

    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            "风险已围绕环境、范围、实时消息和答辩材料四类核心问题进行登记。当前的控制策略是先交付稳定可演示版本，再把复杂图示和增强功能作为明确的后续补项。",
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            {"type": "bullet", "text": "业务真实：是，围绕校园闲置交易与近场服务场景展开。"},
            {"type": "bullet", "text": "规模合适：是，核心范围控制在课程设计可完成的模块化单体系统内。"},
            {"type": "bullet", "text": "技术可行：是，现有技术栈与部署方式均已验证可运行。"},
            {"type": "bullet", "text": "团队胜任：基本具备，建议尽快补齐图示和最终答辩材料。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-01_项目开题报告.docx")


def build_d02() -> None:
    doc = Document(TEMPLATE_DIR / "D-02_需求规格说明书_SRS_模板.docx")
    fill_cover(doc.tables[0], code="D-02", name="软件需求规格说明书 SRS")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "本文档用于明确 SwapCampus 当前课程设计版本的业务边界、用户角色、核心功能、非功能要求与数据口径，为后续概要设计、详细设计、数据库设计、测试和验收提供统一依据。面向读者包括项目组成员、指导教师、评审教师以及后续接手维护的开发者。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "SwapCampus 面向校内师生，提供闲置商品交易、校园服务、站内沟通和后台治理能力。项目旨在替代以微信群和熟人转发为主的碎片化交易方式，在校内身份可核验、商品类型集中、线下面交便利的背景下建立更可信、更高效的校园二手交易体验。",
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            {"type": "bullet", "text": "SRS：Software Requirements Specification，软件需求规格说明书。"},
            {"type": "bullet", "text": "MVP：Minimum Viable Product，最小可行产品。"},
            {"type": "bullet", "text": "LLM：Large Language Model，本项目用于内容审核辅助。"},
            {"type": "bullet", "text": "IM：Instant Messaging，即时通讯能力；本项目当前仅实现演示级实时消息。"},
            {"type": "bullet", "text": "REQUEST / OFFER：校园服务模块中的“求助型发布 / 提供型发布”双向意图。"},
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "bullet", "text": "课程设计任务书与 T-02 选题说明。"},
            {"type": "bullet", "text": "仓库根目录 README.md。"},
            {"type": "bullet", "text": "docs/01-09 主文档、docs/14 验收记录、docs/17 追踪矩阵。"},
            {"type": "bullet", "text": "backend 与 frontend 当前源码实现。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            "需求基线以当前源码、README、测试记录和课程文档为准；若运行时数据与文档示例不一致，以“默认初始化空库建表 + 可选演示种子重建”双口径解释。所有运行中的商品、订单、消息和校园服务数据都被视为动态数据，不作为固定常量写入需求基线。",
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            "产品定位为“仅服务校内师生的校园闲置交易平台”，同时兼容校园服务发布、消息沟通、信用积分和后台治理。它不是泛社会化二手平台，而是以校内近场交易、身份约束和治理留痕为核心价值的课程设计项目。",
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            {"type": "bullet", "text": "账号注册登录与封禁拦截。"},
            {"type": "bullet", "text": "商品浏览、搜索、详情、发布、收藏、审核。"},
            {"type": "bullet", "text": "订单创建、快照、完成、评价、申诉。"},
            {"type": "bullet", "text": "消息会话、文本与附件发送、演示级实时推送。"},
            {"type": "bullet", "text": "校园服务发布、接单、完成、收藏。"},
            {"type": "bullet", "text": "信用中心、关注关系、公开主页。"},
            {"type": "bullet", "text": "后台治理、举报处理、实名审核与审计日志。"},
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "游客：低频浏览，主要关注商品列表、详情和公开信息。"},
            {"type": "bullet", "text": "普通用户：高频使用浏览、发布、收藏、下单、消息和校园服务能力，对操作效率与可信度敏感。"},
            {"type": "bullet", "text": "管理员：中频使用审核、举报、申诉和用户治理能力，关注留痕、状态一致性和处理效率。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            {"type": "bullet", "text": "硬件：普通笔记本或台式机即可运行课程演示环境。"},
            {"type": "bullet", "text": "操作系统：macOS / Linux / Windows（通过 Docker Desktop）。"},
            {"type": "bullet", "text": "浏览器：Chrome、Edge 等现代浏览器。"},
            {"type": "bullet", "text": "网络：本地开发可使用回环地址访问，跨设备演示时需保证局域网互通。"},
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            {"type": "bullet", "text": "课程周期有限，优先保障核心链路稳定。"},
            {"type": "bullet", "text": "不接入真实支付与物流，订单以线下面交完成。"},
            {"type": "bullet", "text": "当前自动化测试以后端 Jest 与前端构建校验为主。"},
            {"type": "bullet", "text": "部分图示、曲线和正式答辩材料仍需在终版前补充。"},
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            {"type": "bullet", "text": "假设校园用户具备学号或邮箱注册条件。"},
            {"type": "bullet", "text": "依赖 Docker Desktop、MySQL、Meilisearch、MinIO 和 SuperTokens 容器可正常启动。"},
            {"type": "bullet", "text": "依赖浏览器支持本地文件上传和基本 WebSocket 能力。"},
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            "系统依赖主要集中在本地容器、对象存储、搜索和认证会话服务。当前不依赖外部支付、地图或短信接口，因此对外部网络环境要求相对可控。",
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            "本轮先以文字化用例关系说明替代正式用例图，建议在答辩终稿中补 1 张总用例图和 1-2 张子用例图。",
            {"type": "bullet", "text": "游客：浏览首页、查看详情、访问公开主页。"},
            {"type": "bullet", "text": "普通用户：注册登录、发布商品、收藏、联系、下单、评价、申诉、发布校园服务。"},
            {"type": "bullet", "text": "管理员：审核商品、处理举报、审核实名、处理申诉、封禁/解封用户。"},
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            {"type": "number", "index": 1, "text": "用例：用户注册与登录。参与者：普通用户。前置：未登录。主流程：输入学号/邮箱与密码，完成注册或登录，系统创建会话并返回用户基础信息。异常：学号重复、邮箱重复、账号封禁、图片上传失败。"},
            {"type": "number", "index": 2, "text": "用例：商品发布与审核。参与者：普通用户、管理员。前置：用户已登录。主流程：上传图片、填写标题价格描述、提交发布；系统进行规则校验与 LLM 审核；管理员可在后台复核状态。异常：图片缺失、内容不合规、权限不足。"},
            {"type": "number", "index": 3, "text": "用例：订单创建与完成。参与者：买家、卖家。前置：商品在售且用户已登录。主流程：买家从详情页创建订单，系统保留商品快照并建立会话，双方线下面交后完成订单，买家可评价。异常：商品下架、订单取消、权限错误。"},
            {"type": "number", "index": 4, "text": "用例：校园服务发布与接单。参与者：普通用户。前置：已登录。主流程：用户以 REQUEST 或 OFFER 方式发布服务，其他用户可收藏、接单、完成或取消。"},
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            "关键用例围绕“账号—商品—订单—消息—治理”主链路组织，当前实现已覆盖核心业务流程，未完成部分主要集中在更完整的 IM 能力和更细粒度的异常处置体验。",
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            {"type": "bullet", "text": "FR-01 用户体系：支持注册、登录、实名补充信息、资料维护、封禁拦截；优先级 Must；来源：课程交易闭环要求。"},
            {"type": "bullet", "text": "FR-02 商品模块：支持浏览、搜索、详情、发布、多图上传、收藏、审核；优先级 Must。"},
            {"type": "bullet", "text": "FR-03 订单模块：支持创建、状态流转、评价、申诉、快照查看；优先级 Must。"},
            {"type": "bullet", "text": "FR-04 消息模块：支持会话列表、明细、附件发送、演示级实时推送；优先级 Should。"},
            {"type": "bullet", "text": "FR-05 校园服务模块：支持 REQUEST / OFFER、接单、完成、收藏；优先级 Should。"},
            {"type": "bullet", "text": "FR-06 后台治理：支持审核、举报、封禁、实名审核、申诉处理、审计日志；优先级 Must。"},
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            {"type": "bullet", "text": "用户体系：输入注册信息与图片，处理校验、同步会话、写入用户与实名数据，输出登录态与用户信息；异常包括重复账号、封禁、上传失败。"},
            {"type": "bullet", "text": "商品发布：输入标题、价格、成色、描述与图片，处理规则校验、媒体上传、审核与入库，输出商品详情页可见数据；异常包括图片缺失、内容违规。"},
            {"type": "bullet", "text": "订单创建：输入商品上下文和买家身份，处理状态校验、生成快照、创建会话，输出订单记录与消息上下文。"},
            {"type": "bullet", "text": "后台审核：输入待处理商品/举报/实名记录，处理状态流转与审计日志，输出最新治理结果。"},
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            "功能需求遵循“当前源码已实现的能力优先入档”的原则，未落地部分仅作为边界说明，不写成已完成需求。",
        ],
    )

    nft = doc.tables[20]
    set_table_row(nft, 1, ["性能", "核心接口响应", "P95 RT", "≤ 350 ms（本机演示目标）"])
    set_table_row(nft, 2, ["可用性", "部署复现成功率", "1 小时内完成部署", "≥ 1 次独立复现成功"])
    set_table_row(nft, 3, ["安全性", "封禁与权限校验", "关键越权用例", "不得绕过"])
    set_table_row(nft, 4, ["可维护性", "前后端构建/测试", "CI 与命令入口", "可执行"])
    extra = nft.add_row()
    for idx, value in enumerate(["可追溯性", "文档与实现同步", "需求-实现-证据链", "主链路均可追溯"]):
        clear_cell(extra.cells[idx])
        write_paragraph(extra.cells[idx].paragraphs[0], value, first_line_indent=False)

    fill_text_cell(
        doc.tables[21].cell(0, 0),
        [
            "当前更偏向课程设计可交付与可复现演示，而非生产环境指标。对性能、安全和稳定性采取“可运行、可验证、可解释”的实现目标。",
        ],
    )
    fill_text_cell(
        doc.tables[22].cell(0, 0),
        [
            {"type": "bullet", "text": "关键页面包括：首页、登录页、商品详情页、发布页、消息页、个人中心和后台页。"},
            {"type": "bullet", "text": "当前交付阶段已形成真实运行截图，详见 D-07 用户手册与 docs/14 验收记录。"},
            {"type": "bullet", "text": "正式答辩建议再补 2-3 张页面线框图，以强化“需求到设计”的过程表达。"},
        ],
    )
    fill_text_cell(
        doc.tables[23].cell(0, 0),
        [
            "N/A。当前项目不接入专用硬件设备，所有功能均运行在浏览器、应用服务和容器化基础设施之上。",
        ],
    )
    fill_text_cell(
        doc.tables[24].cell(0, 0),
        [
            {"type": "bullet", "text": "SuperTokens：认证会话与登录状态管理；协议 HTTP；降级方式为本地开发环境重启认证容器。"},
            {"type": "bullet", "text": "Meilisearch：商品与服务搜索；协议 HTTP；降级方式为保留列表读取但弱化搜索体验。"},
            {"type": "bullet", "text": "MinIO：图片与附件对象存储；协议 S3 兼容 API；降级方式为阻止需要媒体写入的流程。"},
            {"type": "bullet", "text": "Socket.IO：消息实时追加；协议 WebSocket/HTTP fallback；降级方式为退回手动刷新读取消息。"},
        ],
    )
    fill_text_cell(
        doc.tables[25].cell(0, 0),
        [
            {"type": "bullet", "text": "前后端主要通过 REST API 通信。"},
            {"type": "bullet", "text": "消息页通过 Socket.IO 建立会话房间订阅。"},
            {"type": "bullet", "text": "对外接口统一走 HTTP，本地部署通过 127.0.0.1 暴露端口。"},
        ],
    )
    fill_text_cell(
        doc.tables[26].cell(0, 0),
        [
            "通信接口设计遵循简单、可调试、可独立验证原则；当前不引入复杂消息队列或外部事件总线，以控制课程项目复杂度。",
        ],
    )
    fill_text_cell(
        doc.tables[27].cell(0, 0),
        [
            {"type": "bullet", "text": "默认初始化为空库建表，不预置演示账号与业务数据。"},
            {"type": "bullet", "text": "可选演示种子可重建 10 个用户、30 个商品、18 个校园服务、若干订单和举报记录。"},
            {"type": "bullet", "text": "业务数据需保留交易、评价、申诉与审计留痕，以支撑教学验收和问题追溯。"},
            {"type": "bullet", "text": "数据备份与恢复在本地演示环境下以容器卷重建和脚本重置为主。"},
        ],
    )
    fill_text_cell(
        doc.tables[28].cell(0, 0),
        [
            {"type": "bullet", "text": "待补一张正式总用例图与两张子用例图。负责人：待补充。截止：答辩前。"},
            {"type": "bullet", "text": "待补正式覆盖率截图、性能曲线与安全扫描输出。负责人：待补充。截止：终版文档前。"},
            {"type": "bullet", "text": "待补指导教师、组号、成员信息等封面元数据。负责人：待补充。截止：提交前。"},
        ],
    )
    fill_text_cell(
        doc.tables[29].cell(0, 0),
        [
            {"type": "bullet", "text": "评审时间：2026-06-15（本轮文档整理基线）。"},
            {"type": "bullet", "text": "与会人：项目组成员、课程评审教师（待补充实名）。"},
            {"type": "bullet", "text": "主要决议：以当前源码为唯一事实源，文档按模板重新整理，未落地图示单独列清单补齐。"},
            {"type": "bullet", "text": "行动项：生成正式 docx、补架构图/ER 图/时序图、核对封面元数据。"},
        ],
    )
    fill_text_cell(
        doc.tables[30].cell(0, 0),
        [
            "本 SRS 已根据当前源码、测试记录和验收材料重写为可追溯版本，可直接作为 D-03 到 D-06 的需求依据。剩余工作主要是把需要视觉表达的图示补齐到正式终版。",
        ],
    )

    doc.save(OUTPUT_DIR / "D-02_需求规格说明书_SRS.docx")


def build_d03() -> None:
    doc = Document(TEMPLATE_DIR / "D-03_概要设计说明书_模板.docx")
    fill_cover(doc.tables[0], code="D-03", name="概要设计说明书")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "本文档用于说明 SwapCampus 当前版本的总体架构、模块划分、运行视图、部署视图和关键架构决策。主要读者为项目组成员、评审教师以及后续接手维护的开发者。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "本设计分别追溯到 SRS 中的用户体系、商品能力、订单能力、消息能力、校园服务能力、后台治理能力与非功能需求章节，重点响应 D-02 中 2.2、3.2、4.1、4.2、5 和 6 章要求。",
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            "概要设计以“模块化单体 + 可复现基础设施 + 文档与代码同仓”为核心原则，在课程规模内优先追求实现清晰、部署稳定、链路可演示和文档可追溯。",
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "bullet", "text": "可演进：在单体结构内按领域模块拆分，保留后续进一步服务化的可能性。"},
            {"type": "bullet", "text": "可测试：通过清晰的控制器/服务/数据层边界降低联调复杂度。"},
            {"type": "bullet", "text": "可观测：关键治理动作、状态变更和后台处理保留日志与审计信息。"},
            {"type": "bullet", "text": "可部署：统一由 Docker Compose 描述环境，确保教师或组员可重复启动。"},
            {"type": "bullet", "text": "安全：采用会话鉴权、账号封禁拦截、后台权限边界与基础输入校验。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            "当前先以文字化 C4 L1/L2 结构说明替代正式架构图。建议终版补一张“浏览器 / 前端 / backend / search-indexer / commerce-sync / MySQL / SuperTokens / MinIO / Meilisearch / Vendure”总图。",
            {"type": "bullet", "text": "浏览器访问 React 前端。"},
            {"type": "bullet", "text": "前端通过 REST API 与 Socket.IO 连接 NestJS 后端。"},
            {"type": "bullet", "text": "backend 通过 Prisma 访问 MySQL，通过 SuperTokens 管理认证会话，并在主事务内写入 OutboxEvent。"},
            {"type": "bullet", "text": "search-indexer 独立消费 `search.index` 主题，把商品与卖家搜索索引收敛到 Meilisearch。"},
            {"type": "bullet", "text": "commerce-sync 独立消费 `commerce.sync` 主题，把商品、客户和订单异步同步到 Vendure。"},
            {"type": "bullet", "text": "MinIO 提供对象存储，Meilisearch 提供搜索能力，Vendure 提供外围电商同步能力。"},
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            {"type": "bullet", "text": "前端：React + Vite，开发快、组件化清晰、适合课程项目快速迭代。"},
            {"type": "bullet", "text": "后端：NestJS，天然支持模块化、DTO、依赖注入和中大型课程项目表达。"},
            {"type": "bullet", "text": "数据库：MySQL 8 作为业务主库，关系模型清晰；Prisma 提供结构定义与类型支持。"},
            {"type": "bullet", "text": "认证：SuperTokens Session 负责注册登录后的会话生命周期，降低自建鉴权复杂度。"},
            {"type": "bullet", "text": "异步副作用：基于 MySQL Outbox 派生 `search-indexer` 与 `commerce-sync` 两个独立工作进程。"},
            {"type": "bullet", "text": "部署：Docker Compose 将数据库、认证、搜索、对象存储与应用统一编排。"},
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            "当前系统采用前后端分离、主业务模块化单体加独立异步工作进程的总体设计，既能在课程规模内保持主链路集中，又能把搜索索引和外围电商同步这类副作用从请求链路中拆开。",
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "认证与用户模块：负责注册登录、会话同步、资料维护、实名与封禁状态。"},
            {"type": "bullet", "text": "商品与收藏模块：负责商品流、详情、发布、多图、收藏和行为留痕。"},
            {"type": "bullet", "text": "订单与评价模块：负责下单、状态流转、评价、申诉与快照。"},
            {"type": "bullet", "text": "消息模块：负责会话、消息附件与 Socket.IO 演示级实时推送。"},
            {"type": "bullet", "text": "校园服务模块：负责 REQUEST / OFFER 双向服务发布与接单流程。"},
            {"type": "bullet", "text": "后台治理模块：负责审核、举报、申诉、实名、封禁与审计。"},
            {"type": "bullet", "text": "基础设施模块：负责 OutboxEvent、搜索索引异步消费和 Vendure 异步同步。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            "模块依赖以“前端页面 -> 后端控制器 -> 领域服务 -> Prisma/外部组件”为主。治理模块依赖商品、订单、服务、用户等领域数据；消息模块与订单、商品、校园服务通过会话上下文关联；搜索索引和 Vendure 同步则通过 Outbox 事件把副作用从主事务中解耦。",
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            "当前模块关系清晰，但正式终稿仍建议补一张依赖图，突出认证、媒体、搜索、交易和治理之间的调用边界。",
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            {"type": "number", "index": 1, "text": "登录流程：用户提交学号/邮箱与密码 -> 后端校验账号状态 -> SuperTokens 建立会话 -> 前端写入当前登录态并跳转。"},
            {"type": "number", "index": 2, "text": "下单流程：买家从商品详情创建订单 -> 后端校验商品状态 -> 创建订单与商品快照 -> 生成或绑定会话 -> 前端进入订单与消息上下文。"},
            {"type": "number", "index": 3, "text": "校园服务流程：用户发布 REQUEST/OFFER -> 其他用户接单 -> 双方确认完成 -> 状态写回与消息通知。"},
            "建议在终稿补 2 张正式时序图，优先补“登录流程”和“下单流程”。",
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            {"type": "bullet", "text": "订单状态：PENDING -> IN_PROGRESS -> WAITING_REVIEW -> COMPLETED；异常分支可到 CANCELED。"},
            {"type": "bullet", "text": "商品状态：PENDING -> ON_SALE -> SOLD / OFFLINE。"},
            {"type": "bullet", "text": "校园服务状态：OPEN -> BUSY / PAUSED -> ENDED 或 CANCELED。"},
            {"type": "bullet", "text": "账号状态：ACTIVE / BANNED。实名状态：PENDING / APPROVED / REJECTED。"},
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            "状态机已在数据库枚举与服务逻辑中落地，正式答辩材料建议补 1 张订单状态机图和 1 张校园服务状态机图。",
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            "部署视图采用 Docker Compose 拓扑：`frontend`（nginx 静态托管）、`backend`（业务主服务）、`search-indexer` 与 `commerce-sync` 作为应用层，`mysql`、`supertokens-db`、`supertokens`、`minio`、`meilisearch`、`vendure` 作为基础能力层，`db-init` 作为一次性初始化任务。当前建议单独补一张部署拓扑图。",
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            "高层数据视图由用户、商品、订单、消息、校园服务、举报与信用中心几大实体组成。数据关系详见 D-05；正式终稿建议补 1 张全局 ER 图和 1 张交易/消息上下文子图。",
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            "Status：Accepted。Context：课程周期有限，但系统能力涉及商品、订单、消息、校园服务、治理与部署，同时搜索和外围电商同步又不适合继续耦合在主请求链路中。Decision：主业务保持 NestJS 模块化单体，同时把搜索索引与 Vendure 同步抽成独立异步工作进程，而不是把核心领域提前拆成多个自治服务。Consequences：主链路实现集中、部署仍可控，且能体现事件驱动解耦；代价是需要额外维护 Outbox、轮询消费和失败补偿。Alternatives：一开始就全面微服务化复杂度过高，不适合当前阶段。",
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            "Status：Accepted。Context：业务主数据强依赖关系模型，且团队使用 Prisma 进行结构管理。Decision：业务主库采用 MySQL 8；认证侧遵循 SuperTokens 默认 PostgreSQL 拆分。Consequences：业务表建模与查询成本较低，但需要注意双库边界。Alternatives：纯 PostgreSQL 统一建模在本轮不是必要条件。",
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            "Status：Accepted。Context：需要支持注册登录、封禁拦截和课程演示环境下稳定的会话维持。Decision：采用 SuperTokens Session 而非自建 JWT 链路。Consequences：降低鉴权实现成本，提高会话可靠性，但需要维护业务用户表与认证用户的同步。Alternatives：自建 JWT 需要额外处理 refresh、rotation、登出与黑名单问题。",
        ],
    )
    fill_text_cell(
        doc.tables[20].cell(0, 0),
        [
            "三项 ADR 已覆盖架构边界、数据库方案与鉴权方案，能够解释当前课程项目为什么选择“主业务模块化单体 + MySQL Outbox + 独立异步 worker + SuperTokens 会话”的总体路线。",
        ],
    )
    fill_text_cell(
        doc.tables[21].cell(0, 0),
        [
            {"type": "bullet", "text": "对外接口以 `/api/*` REST 路径为主，关键端点覆盖认证、商品、订单、消息、校园服务、后台治理与信用中心。"},
            {"type": "bullet", "text": "鉴权方式以 Session 为主，普通接口按登录态控制，后台接口按管理员权限控制。"},
            {"type": "bullet", "text": "错误码与错误信息遵循 NestJS 异常返回模式，重点处理权限不足、状态不合法和资源不存在。"},
        ],
    )
    fill_text_cell(
        doc.tables[22].cell(0, 0),
        [
            {"type": "bullet", "text": "内部接口主要表现为模块间服务调用与共享数据契约。"},
            {"type": "bullet", "text": "消息模块通过会话上下文与订单、商品、校园服务建立关联，不额外引入独立 RPC。"},
            {"type": "bullet", "text": "媒体模块统一承接图片和附件上传，向商品、实名、消息等模块提供公共能力。"},
        ],
    )
    fill_text_cell(
        doc.tables[23].cell(0, 0),
        [
            "当前接口设计已满足课程演示与部署验证要求；若正式提交需要更强过程性表达，建议补充 OpenAPI 摘要截图或导出文档。",
        ],
    )
    fill_text_cell(
        doc.tables[24].cell(0, 0),
        [
            {"type": "bullet", "text": "错误分为参数校验、权限校验、状态冲突、资源缺失和外部依赖失败五类。"},
            {"type": "bullet", "text": "日志以容器日志、后端应用日志和审计记录为主，关键治理动作保留可追溯留痕。"},
            {"type": "bullet", "text": "当前未全面引入 traceId 体系，正式终稿可将其列为后续演进点。"},
        ],
    )
    fill_text_cell(
        doc.tables[25].cell(0, 0),
        [
            {"type": "bullet", "text": "认证：基于 SuperTokens Session 维持登录态。"},
            {"type": "bullet", "text": "鉴别：普通用户与管理员在后台接口层做显式区分。"},
            {"type": "bullet", "text": "输入与媒体安全：上传、发布与消息写入前做基础校验与审核。"},
            {"type": "bullet", "text": "审计：举报处理、申诉处理、审核与封禁等动作保留审计视角。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-03_概要设计说明书.docx")


def build_d04() -> None:
    doc = Document(TEMPLATE_DIR / "D-04_详细设计说明书_模板.docx")
    fill_cover(doc.tables[0], code="D-04", name="详细设计说明书")

    set_heading(doc.paragraphs[26], "3.1  认证与用户模块")
    set_heading(doc.paragraphs[28], "3.2  商品与订单模块")
    set_heading(doc.paragraphs[30], "3.3  消息与校园服务模块")
    set_heading(doc.paragraphs[34], "4.1  状态流转与幂等控制")
    set_heading(doc.paragraphs[36], "4.2  推荐与排序信号利用")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "本文档用于把概要设计中的模块划分进一步展开为实现级说明，回答每个核心模块的职责、接口、关键对象、状态流转和异常处理策略，为课程项目的编码、联调和测试提供实现依据。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "本说明书承接 D-03 中的模块化单体设计思路，对认证与用户、商品与订单、消息与校园服务三个核心模块做重点展开，并补充公共数据结构、工具能力、算法边界与测试设计要点。",
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            "详细设计聚焦于当前已经落地的真实实现，不额外虚构未实现的类或算法；对于尚未图示化的类图和时序图，在相应章节以文字说明并纳入待补图清单。",
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "bullet", "text": "前后端统一采用 TypeScript。"},
            {"type": "bullet", "text": "后端遵循 NestJS 模块/控制器/服务分层，使用 DTO 做参数校验。"},
            {"type": "bullet", "text": "前端按页面、组件、服务和类型分层组织。"},
            {"type": "bullet", "text": "统一通过 Lint、构建和测试命令做基础质量校验。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            {"type": "bullet", "text": "统一分页与列表响应结构。"},
            {"type": "bullet", "text": "统一错误响应依赖 NestJS HttpException 体系。"},
            {"type": "bullet", "text": "统一状态枚举包括商品状态、订单状态、消息类型、实名状态、账号状态等。"},
            {"type": "bullet", "text": "统一异步状态包括 `ExternalSyncStatus` 与 `OutboxEventStatus`。"},
            {"type": "bullet", "text": "统一会话上下文通过 Conversation/Message 模型关联。"},
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            {"type": "bullet", "text": "Prisma Client：统一数据库访问。"},
            {"type": "bullet", "text": "媒体上传工具：统一处理商品图片、头像、学生证和消息附件。"},
            {"type": "bullet", "text": "认证同步服务：统一处理业务用户表与 SuperTokens 之间的数据同步。"},
            {"type": "bullet", "text": "OutboxService：统一处理搜索索引与 Vendure 同步事件的发布。"},
            {"type": "bullet", "text": "消息网关服务：统一处理 Socket.IO 房间订阅与消息推送。"},
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            "全局约定以“结构定义统一、状态流转清晰、跨模块上下文复用、主事务副作用异步解耦”为重点，减少课程项目中常见的接口漂移和状态失配问题。",
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "职责：注册登录、资料同步、实名信息维护、封禁校验、关注与公开主页。"},
            {"type": "bullet", "text": "关键接口：`POST /api/auth/register`、`POST /api/auth/login`、`GET /api/auth/me`、`PATCH /api/users/:id/profile`。"},
            {"type": "bullet", "text": "关键对象：User、StudentVerification、UserFollow、Session 上下文。"},
            {"type": "bullet", "text": "实现要点：注册支持头像与学生证图片上传；登录时同时校验账号状态；管理员可修改封禁和实名状态。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            {"type": "bullet", "text": "职责：商品列表、详情、搜索、发布、多图上传、收藏、下单、评价、申诉。"},
            {"type": "bullet", "text": "关键接口：`GET /api/products`、`POST /api/products`、`POST /api/orders`、`POST /api/orders/:id/reviews`、`POST /api/orders/:id/appeals`。"},
            {"type": "bullet", "text": "关键对象：Product、ProductImage、Favorite、Order、Review、OrderAppeal、UserBehavior。"},
            {"type": "bullet", "text": "实现要点：发布流程先上传图片，再写入商品数据，并在事务内发布 `search.index` 与 `commerce.sync` 事件；订单会保留商品快照、自动建立消息上下文，并在创建、取消、完成时同步更新本地异步同步状态。"},
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            {"type": "bullet", "text": "职责：会话列表、消息明细、文本与附件发送、Socket.IO 房间推送、校园服务发布与接单。"},
            {"type": "bullet", "text": "关键接口：`GET /api/messages/conversations`、`POST /api/messages/conversations/:id`、`GET /api/campus-services`、`POST /api/campus-services/:id/orders`。"},
            {"type": "bullet", "text": "关键对象：Conversation、Message、CampusServiceListing、CampusServiceOrder、CampusServiceFavorite。"},
            {"type": "bullet", "text": "实现要点：REST 承担读写，Socket.IO 负责演示级实时追加；校园服务采用 listing + order 双层建模支持 REQUEST/OFFER 双向语义。"},
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            "三个核心模块覆盖了当前课程项目的绝大多数业务复杂度，剩余后台治理、信用中心和媒体模块更多是对上述主链路的扩展与支撑；搜索索引和 Vendure 同步则由独立 worker 在主链路之外异步处理。",
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            {"type": "bullet", "text": "订单与校园服务状态流转均采用显式状态枚举与前置状态校验。"},
            {"type": "bullet", "text": "所有状态迁移在服务层完成，避免由前端任意构造。"},
            {"type": "bullet", "text": "创建订单与发送消息时优先检查上下文是否已存在，减少重复会话或非法推进。"},
            {"type": "bullet", "text": "搜索索引与 Vendure 同步通过 OutboxEvent + 独立消费者处理，消费状态与业务同步状态分开记录。"},
            {"type": "bullet", "text": "复杂度以状态数量和迁移边数线性增长，适合当前项目规模。"},
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            {"type": "bullet", "text": "首页推荐与排序主要利用分类、时间、收藏/浏览等行为信号。"},
            {"type": "bullet", "text": "当前实现偏启发式排序，而非复杂机器学习推荐。"},
            {"type": "bullet", "text": "推荐逻辑优先保证可解释性和课程演示效果，后续可继续接入更细粒度的行为权重。"},
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            "当前项目的关键算法重在状态控制、规则校验和启发式排序，而非复杂数学模型，因此文档也以业务算法说明为主。",
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            "当前仓库未单独维护正式 UML 类图。建议终稿补 1 张用户/商品/订单核心域类图和 1 张消息/校园服务上下文类图。现有文字版核心对象已见本章 3.1-3.3。",
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            "建议终稿至少补 2 张时序图：一张覆盖“注册登录并进入首页”，一张覆盖“商品下单并创建会话”。当前文字化时序已分别在 D-03 和本章中描述。",
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            "类图和时序图是本说明书当前最明显的可视化缺口，已在单独清单中列出，便于后续集中补图。",
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            {"type": "bullet", "text": "异常处理：对参数错误、资源不存在、权限不足和状态冲突做显式异常返回。"},
            {"type": "bullet", "text": "事务边界：订单、申诉、积分等关键写入优先在同一业务流程内保证一致性；业务数据变更与 OutboxEvent 写入保持同事务提交。"},
            {"type": "bullet", "text": "并发策略：当前部署规模较小，主要通过数据库唯一约束、状态检查和幂等式写法规避重复提交。"},
            {"type": "bullet", "text": "消息一致性：先写数据库再广播，避免前端收到未落库消息。"},
            {"type": "bullet", "text": "异步补偿：独立 worker 通过 `PENDING / PROCESSING / PROCESSED / FAILED` 或同步状态字段执行重试与失败留痕。"},
        ],
    )
    fill_text_cell(
        doc.tables[20].cell(0, 0),
        [
            {"type": "bullet", "text": "认证模块重点覆盖注册、登录、封禁拦截、资料同步和实名审核边界。"},
            {"type": "bullet", "text": "商品与订单模块重点覆盖发布校验、收藏切换、状态流转、评价与申诉。"},
            {"type": "bullet", "text": "消息与校园服务模块重点覆盖会话权限、附件写入、实时追加和 REQUEST/OFFER 状态分支。"},
            {"type": "bullet", "text": "异步链路重点覆盖 Outbox 事件发布、worker 消费、失败重试和状态回写。"},
            {"type": "bullet", "text": "外部依赖如对象存储、认证服务和搜索组件建议通过 Mock 或容器联调双方式验证。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-04_详细设计说明书.docx")


def build_d05() -> None:
    doc = Document(TEMPLATE_DIR / "D-05_数据库设计说明书_模板.docx")
    fill_cover(doc.tables[0], code="D-05", name="数据库设计说明书")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "本文档用于说明 SwapCampus 当前版本的数据组织方式、主实体关系、逻辑建模、索引约束、安全策略与数据生命周期，支撑交易、消息、校园服务和治理等核心业务链路。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "业务主数据库采用 MySQL 8；认证会话依赖 SuperTokens 独立 PostgreSQL。选择理由包括：交易与治理数据以结构化关系为主、Prisma 与 MySQL 配合成熟、课程项目需要快速稳定迭代，同时将认证数据与业务数据适度分离可以减少强耦合。",
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            "当前数据库设计由 `backend/prisma/schema.prisma` 作为唯一结构事实源，配合 `make init` 和 Prisma 命令完成建表。正式答辩建议补 1 张全局 ER 图和 2 张交易/消息/校园服务子图。",
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            "概念模型可分为六组实体：用户与实名、商品与收藏、订单与评价申诉、消息会话、校园服务、信用与治理。全局上以 User 为核心节点，向商品、订单、校园服务、消息、关注和信用资产延伸。",
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            {"type": "bullet", "text": "User(studentId, email, role, accountStatus, verificationStatus, ... )。"},
            {"type": "bullet", "text": "Product(id, sellerId, category, status, price, ... )；ProductImage(id, productId, url, ... )。"},
            {"type": "bullet", "text": "Favorite(userId, productId)；UserFollow(followerId, followingId)。"},
            {"type": "bullet", "text": "Order(id, productId, buyerId, sellerId, status, snapshot, ... )；Review(orderId, authorId, ... )；OrderAppeal(orderId, appellantId, ... )。"},
            {"type": "bullet", "text": "Conversation(id, contextType, contextId, ... )；Message(id, conversationId, senderId, type, ... )。"},
            {"type": "bullet", "text": "CampusServiceListing(id, ownerId, intent, status, ... )；CampusServiceOrder(id, listingId, requesterId, providerId, status, ... )。"},
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            {"type": "bullet", "text": "整体逻辑模型按第三范式设计，避免把多值属性直接堆叠在单表中。"},
            {"type": "bullet", "text": "商品图片、消息附件、校园服务图片均拆分为独立对象或媒体记录。"},
            {"type": "bullet", "text": "订单保留商品快照属于面向交易追溯的受控反范式设计，用于避免商品后续修改影响历史订单解释。"},
            {"type": "bullet", "text": "信用资产采用聚合表 + 流水表组合，兼顾读取效率与留痕完整性。"},
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            "逻辑模型重点保证“一个事实只在一个核心表中维护”，对于快照、统计和推荐等读优化数据则采用有限反范式，以换取交易解释性和查询效率。",
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "当前建议优先整理 User、Product、Order、Conversation、CampusServiceListing 五张核心表的数据字典。"},
            {"type": "bullet", "text": "敏感字段包括邮箱、手机号、实名姓名、学生证图片地址等，需要在展示层与权限层做区分。"},
            {"type": "bullet", "text": "完整字段定义以 `backend/prisma/schema.prisma` 为准，当前文档强调结构意图与关系设计。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            "当前项目以 Prisma schema 作为结构真源，不手工维护整份 CREATE TABLE 脚本。建表方式为 `npm run db:push -- --force-reset` 或 `make init`；正式归档时可从实际数据库导出 SQL 作为附录补充。",
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            "物理模型与 DDL 在仓库中采用“schema 驱动生成”而非“手写 SQL 主维护”的工程方式，这更符合课程项目快速演进与类型一致性的实际需要。",
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            {"type": "bullet", "text": "唯一约束：User.studentId、User.email、Favorite(userId, productId)、CampusServiceFavorite(userId, listingId)、UserFollow(followerId, followingId)。"},
            {"type": "bullet", "text": "常用索引：Product 按 sellerId/status/category；Order 按 buyerId/sellerId/status；OrderAppeal 按 orderId/status；行为表按 userId + createdAt。"},
            {"type": "bullet", "text": "搜索检索更多依赖 Meilisearch，因此数据库索引重点保证事务链路和后台治理查询。"},
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            {"type": "bullet", "text": "当前可选演示种子规模约为：10 用户、30 商品、18 校园服务、若干订单和举报。"},
            {"type": "bullet", "text": "慢查询风险主要来自商品列表筛选、后台多条件查询、消息上下文读取和行为统计。"},
            {"type": "bullet", "text": "治理思路包括：复用联合索引、避免不必要的跨表深联查、把全文检索交给 Meilisearch、对详情类请求采用上下文聚合读取。"},
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            "数据规模在课程阶段仍属中小型，但通过明确索引和结构边界，可以为后续扩展到更大样本量预留空间。",
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            {"type": "bullet", "text": "开发与演示环境下主要通过应用统一连接数据库，不开放给普通用户直接访问。"},
            {"type": "bullet", "text": "数据库权限按照应用服务最小权限原则配置；管理操作通过后台 API 完成，而非手工直连数据库。"},
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            {"type": "bullet", "text": "邮箱、手机号、真实姓名、学生证图片等属于敏感或半敏感字段。"},
            {"type": "bullet", "text": "前端展示时按页面职责决定是否显示完整信息；公开主页仅展示必要公开字段。"},
            {"type": "bullet", "text": "学生证图片与附件统一存放于对象存储，并通过业务权限控制访问。"},
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            {"type": "bullet", "text": "关键治理动作通过审计日志和业务留痕保留可追溯记录。"},
            {"type": "bullet", "text": "备份与恢复在本地演示环境中主要依赖容器卷清理、脚本重建和种子回灌。"},
            {"type": "bullet", "text": "正式答辩可增加一次数据库导出/导入操作截图，增强运维可信度。"},
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            "当前安全与合规策略以课程设计可解释和可复现为主，生产级加密、细粒度脱敏与长期备份策略仍可继续增强。",
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            {"type": "bullet", "text": "默认初始化：只建表，不导入业务种子数据。"},
            {"type": "bullet", "text": "演示数据：通过 `npm run db:reset-and-seed-demo` 重建。"},
            {"type": "bullet", "text": "交易、评价、申诉、举报和审计数据在课程验收期内建议保留，不做自动清理。"},
            {"type": "bullet", "text": "若需要归档或销毁，应优先删除演示数据，再重新执行空库初始化。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-05_数据库设计说明书.docx")


def build_d06() -> None:
    doc = Document(TEMPLATE_DIR / "D-06_软件测试计划与测试报告_模板.docx")
    fill_cover(doc.tables[0], code="D-06", name="软件测试计划与测试报告")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "测试目标是验证 SwapCampus 在当前课程设计版本下具备可运行、可部署、可演示的核心链路，重点覆盖交易、消息、校园服务、后台治理与部署启动能力。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            {"type": "bullet", "text": "包含：注册登录、商品发布与浏览、收藏、订单、消息、校园服务、后台治理、Docker 启动、前端构建与后端测试。"},
            {"type": "bullet", "text": "不包含：真实支付、跨校扩展、完整 IM 在线状态、系统化性能压测和正式安全扫描闭环。"},
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            "测试范围以当前已落地功能为准，未完成或仅处于增强规划阶段的能力不纳入“已通过”结论。",
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "bullet", "text": "单元/服务级验证：以后端 Jest 测试为主。"},
            {"type": "bullet", "text": "集成验证：以前后端联调、真实数据库读写、对象存储上传和后台治理操作为主。"},
            {"type": "bullet", "text": "E2E：当前未系统化接入，只做人工主链路演示验证。"},
            {"type": "bullet", "text": "性能与安全：目前以边界分析和配置核查为主，缺少正式压测曲线和扫描报告。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            {"type": "bullet", "text": "功能测试：接口联调、页面验证、业务流程检查。"},
            {"type": "bullet", "text": "构建验证：`npm run build --prefix frontend`。"},
            {"type": "bullet", "text": "后端测试：`npm test --prefix backend`。"},
            {"type": "bullet", "text": "部署验证：`make init`、`make start` 与健康检查接口。"},
            {"type": "bullet", "text": "兼容性与可访问性：本轮未做系统化专项测试。"},
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            {"type": "bullet", "text": "dev：React 18 + Vite、NestJS + Prisma、本机 Docker Compose。"},
            {"type": "bullet", "text": "数据基线 1：空库建表，不预置演示账号。"},
            {"type": "bullet", "text": "数据基线 2：执行演示种子脚本后形成完整示例数据。"},
            {"type": "bullet", "text": "配置差异：前端 `VITE_*` 和后端 API 域名需随部署地址调整。"},
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            {"type": "bullet", "text": "核心接口与页面可访问。"},
            {"type": "bullet", "text": "商品、订单、消息、校园服务和后台治理主链路均完成至少一次正向验证。"},
            {"type": "bullet", "text": "后端测试命令、前端构建命令和 Docker 启动命令可成功执行。"},
            {"type": "bullet", "text": "若性能、安全、E2E 未执行，应在报告中明确列为边界而非默认通过。"},
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            "退出准则采用“课程交付可运行”口径，而非生产上线口径，因此特别强调可部署、可演示、可追溯，而不夸大未执行的性能/安全结论。",
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            {"type": "bullet", "text": "TC-01 用户注册与登录：通过。"},
            {"type": "bullet", "text": "TC-02 商品图片上传与发布：通过。"},
            {"type": "bullet", "text": "TC-03 登录用户商品收藏：通过。"},
            {"type": "bullet", "text": "TC-04 订单创建与会话联动：通过。"},
            {"type": "bullet", "text": "TC-05 订单评价：通过。"},
            {"type": "bullet", "text": "TC-06 订单申诉：通过。"},
            {"type": "bullet", "text": "TC-07 消息发送与实时追加：通过。"},
            {"type": "bullet", "text": "TC-08 校园服务发布与接单：通过。"},
            {"type": "bullet", "text": "TC-09 Docker 启动、前端构建、后端测试：通过。"},
        ],
    )

    exec_table = doc.tables[11]
    set_table_row(exec_table, 1, ["单元", "12", "12", "12", "0", "0", "100%"])
    set_table_row(exec_table, 2, ["集成", "17", "17", "17", "0", "0", "100%"])
    set_table_row(exec_table, 3, ["E2E", "4", "0", "0", "0", "4", "0%"])
    set_table_row(exec_table, 4, ["性能", "1", "0", "0", "0", "1", "0%"])
    set_table_row(exec_table, 5, ["安全", "1", "0", "0", "0", "1", "0%"])

    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            "已执行部分主要集中在单元、集成、部署和人工业务联调；E2E、性能与安全专项当前尚未系统开展，因此在执行统计中明确标记为阻塞/未执行。",
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            "当前未单独产出覆盖率截图。已执行的自动化验证包括后端 Jest 测试与前端生产构建，建议后续补充 Jest coverage 或 Istanbul 报告截图。",
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            "本轮未进行正式压测，因此无 RT/TPS/错误率曲线。若需要增强本节，可围绕商品列表、搜索和消息读取三类接口做轻量压测并补充曲线截图。",
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            "本轮未执行 ZAP、Trivy 等正式扫描工具。当前安全结论仅限于权限边界、封禁拦截、上传校验和后台治理流程的人工验证，不应替代正式安全扫描。",
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            {"type": "bullet", "text": "当前未形成独立缺陷台账，已发现并关闭的问题主要集中在容器依赖注入、端口冲突、消息归属渲染、实时推送接入和文档口径不一致。"},
            {"type": "bullet", "text": "严重度以中高为主，现阶段均已修复或纳入边界说明。"},
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            {"type": "bullet", "text": "已知风险：完整 IM、E2E、性能与安全专项尚未补齐。"},
            {"type": "bullet", "text": "建议：正式答辩前补 1 轮轻量图示化证据，包括覆盖率、压测或扫描截图。"},
            {"type": "bullet", "text": "上线建议：当前适合作为课程演示与交付版本，不宜按生产系统直接推广。"},
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            {"type": "bullet", "text": "自动化测试代码路径：`backend/` 测试目录与命令脚本。"},
            {"type": "bullet", "text": "部署验证与验收说明：`docs/14-验收与演示记录.md`。"},
            {"type": "bullet", "text": "交付与运行说明：`docs/08-部署与运维手册.md`、`README.md`。"},
            {"type": "bullet", "text": "若后续补性能/安全专项，可将截图或报告文件归档到 `artifacts/`。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-06_软件测试计划与测试报告.docx")


def build_d07() -> None:
    doc = Document(TEMPLATE_DIR / "D-07_用户手册_模板.docx")
    fill_cover(doc.tables[0], code="D-07", name="用户手册")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "SwapCampus 用于帮助校内师生更方便地发布、浏览和交易闲置物品，并在同一系统中完成联系、下单、评价、举报和校园服务协作。对普通用户来说，它比微信群更容易查找信息、更容易保留记录，也更便于老师在演示时看到完整业务闭环。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "最常用的一条上手路径是：打开首页 -> 登录或注册 -> 浏览商品 -> 查看详情 -> 发布或联系卖家。",
            {"type": "image", "path": IMAGE_DIR / "home-final.png", "caption": "图 1 首页与商品流", "width": Cm(14.0)},
            {"type": "image", "path": IMAGE_DIR / "login-final.png", "caption": "图 2 登录与注册页面", "width": Cm(12.0)},
            {"type": "image", "path": IMAGE_DIR / "campus-services-final.png", "caption": "图 3 校园服务页面", "width": Cm(14.0)},
            {"type": "image", "path": IMAGE_DIR / "detail-final.png", "caption": "图 4 商品详情页", "width": Cm(14.0)},
            {"type": "image", "path": IMAGE_DIR / "publish-final.png", "caption": "图 5 商品发布页", "width": Cm(12.5)},
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            {"type": "bullet", "text": "游客：可直接浏览首页、商品详情和公开主页，但不能发布、收藏、下单或发消息。"},
            {"type": "bullet", "text": "普通用户：注册后可使用商品、订单、消息、校园服务、信用中心等完整功能。"},
            {"type": "bullet", "text": "管理员：通过后台页处理商品审核、举报、申诉、实名审核和用户封禁。"},
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "bullet", "text": "登录与注册：从首页右上角进入登录页，支持学号或邮箱登录，也支持新用户注册。"},
            {"type": "bullet", "text": "首页与搜索：在首页浏览推荐商品、分类和活动入口；搜索页支持关键词、价格与信用筛选。"},
            {"type": "bullet", "text": "商品详情与收藏：查看商品图片、价格、卖家和信用摘要；登录后可收藏、联系卖家、下单或举报。"},
            {"type": "bullet", "text": "发布商品：进入 `/publish`，上传至少一张图片并填写标题、价格、成色和描述后提交。"},
            {"type": "bullet", "text": "消息：在消息页查看会话、发送文本或附件，当前支持演示级实时追加。"},
            {"type": "bullet", "text": "订单：从商品详情创建订单，可完成、取消、评价和申诉。"},
            {"type": "bullet", "text": "校园服务：支持 REQUEST / OFFER 两类发布，可收藏、接单、暂停、结束和完成。"},
            {"type": "bullet", "text": "个人中心与信用中心：查看订单、收藏、关注、浏览历史、实名状态、签到和积分任务。"},
            {"type": "bullet", "text": "后台：管理员可在后台页进行审核、举报处理、申诉处理、实名审核和用户治理。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            {"type": "bullet", "text": "Q1：为什么我能浏览商品却不能下单？A：游客模式仅支持浏览，发布、收藏、下单和发消息需要登录。"},
            {"type": "bullet", "text": "Q2：支持什么方式登录？A：支持学号或邮箱登录。"},
            {"type": "bullet", "text": "Q3：发布商品时最少上传几张图？A：至少 1 张。"},
            {"type": "bullet", "text": "Q4：收藏会不会跨设备同步？A：登录用户会随账号同步，游客临时想要只保留在当前浏览器。"},
            {"type": "bullet", "text": "Q5：订单是线上支付吗？A：不是，本项目以线下面交为主。"},
            {"type": "bullet", "text": "Q6：为什么商品提交后看不到？A：可能还在审核中，或内容未通过校验。"},
            {"type": "bullet", "text": "Q7：消息为什么没有完整聊天软件那样的功能？A：当前只实现了课程演示所需的会话和实时追加。"},
            {"type": "bullet", "text": "Q8：校园服务和商品有什么区别？A：商品关注实物交易，校园服务关注委托、代办、跑腿等任务。"},
            {"type": "bullet", "text": "Q9：后台是谁可以看到？A：仅管理员角色可访问。"},
            {"type": "bullet", "text": "Q10：如何恢复演示数据？A：由维护人员执行 `npm run db:reset-and-seed-demo --prefix backend`。"},
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            {"type": "bullet", "text": "“请先登录”：当前操作需要登录账号。"},
            {"type": "bullet", "text": "“学号必须为 9 位数字”：注册或资料维护中的学号格式不正确。"},
            {"type": "bullet", "text": "“邮箱或学号已被占用”：当前注册信息已被使用。"},
            {"type": "bullet", "text": "“账号已被封禁”：账号当前不可登录或执行关键操作。"},
            {"type": "bullet", "text": "“至少上传 1 张图片”：发布商品时图片数量不足。"},
            {"type": "bullet", "text": "“权限不足”或“资源不存在”：可能访问了不属于自己的订单、会话或后台功能。"},
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            {"type": "bullet", "text": "账号注销与数据导出当前未做独立用户界面，若需要处理应由维护人员协助。"},
            {"type": "bullet", "text": "举报功能已支持提交，管理员可在后台处理。"},
            {"type": "bullet", "text": "实名信息、头像和学生证图片仅用于账户与治理相关流程，不在公开页面直接展示。"},
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "课程演示阶段默认通过项目组成员现场反馈与文档问题清单收集。"},
            {"type": "bullet", "text": "仓库说明入口：README.md、D-08 部署与运维手册、D-14 验收与演示记录。"},
            {"type": "bullet", "text": "如需新增正式联系方式，可在提交前补充课程邮箱或团队工单渠道。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-07_用户手册.docx")


def build_d08() -> None:
    doc = Document(TEMPLATE_DIR / "D-08_部署与运维手册_模板.docx")
    fill_cover(doc.tables[0], code="D-08", name="部署与运维手册")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "当前部署拓扑可概括为：浏览器访问 `frontend` 容器（nginx 托管静态资源），前端通过 HTTP 与 Socket.IO 连接 `backend`，`backend` 在主事务中写入 OutboxEvent；`search-indexer` 负责消费搜索事件并访问 `meilisearch`，`commerce-sync` 负责消费电商同步事件并访问 `vendure`，底层统一依赖 `mysql`、`supertokens`、`minio` 等服务。建议在终版补一张正式部署拓扑图。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            {"type": "bullet", "text": "课程演示环境：普通笔记本即可。"},
            {"type": "bullet", "text": "建议配置：4 核 CPU、8GB 内存、20GB 可用磁盘。"},
            {"type": "bullet", "text": "容器数量：默认 10 个运行态容器 + 1 个可选初始化任务容器。"},
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            {"type": "bullet", "text": "必需：Docker Desktop。"},
            {"type": "bullet", "text": "可选：Node.js 20+、npm 10+，用于本机构建、测试和脚本执行。"},
            {"type": "bullet", "text": "数据库与中间件版本以容器镜像定义为准，不需要宿主机单独安装。"},
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            {"type": "bullet", "text": "前端：5178。"},
            {"type": "bullet", "text": "后端 API：3001。"},
            {"type": "bullet", "text": "MySQL：3306。"},
            {"type": "bullet", "text": "Meilisearch：7700。"},
            {"type": "bullet", "text": "MinIO API / Console：9000 / 9001。"},
            {"type": "bullet", "text": "SuperTokens PostgreSQL：5433。"},
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            "当前项目未接入短信、支付或地图等第三方外部业务服务。外部依赖主要是对象存储、搜索、认证和 Vendure 外围电商同步组件，统一由 Docker Compose 内部编排管理。",
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            "环境与依赖已收敛为“Docker 为主，本机 Node 为辅”的低门槛交付方式，便于课程演示与同学之间复现。",
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            {"type": "bullet", "text": "准备根目录 `.env`，根据部署地址调整 `API_DOMAIN`、`WEBSITE_DOMAIN` 与 `VITE_*` 变量。"},
            {"type": "bullet", "text": "确认 Docker Desktop 已启动，且本机 3001、5178、3306、7700、9000、9001、5433 等端口未被冲突占用。"},
            {"type": "bullet", "text": "首次部署前建议执行一次 `docker compose down -v` 清理旧卷（若可接受清空数据）。"},
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "首次启动：`make init`，预期 `db-init` 完成空库建表并退出 0。"},
            {"type": "bullet", "text": "运行应用：`make start`，预期 frontend、backend、search-indexer、commerce-sync、mysql、minio、meilisearch、vendure、supertokens 等容器均为 running。"},
            {"type": "bullet", "text": "健康检查：`curl http://127.0.0.1:3001/api/health`，预期返回 `{\"status\":\"ok\"}`。"},
            {"type": "bullet", "text": "如需演示数据：`npm run db:reset-and-seed-demo --prefix backend`。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            {"type": "bullet", "text": "访问首页 `http://127.0.0.1:5178` 正常显示商品流。"},
            {"type": "bullet", "text": "后端健康检查返回 `ok`。"},
            {"type": "bullet", "text": "消息、校园服务、后台等关键页面可打开并拉取数据。"},
            {"type": "bullet", "text": "若刚执行初始化，`db-init` 状态应为 `exited (0)`；`search-indexer` 与 `commerce-sync` 应处于运行状态。"},
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            "从零部署已经具备明确的命令入口，但正式提交仍建议补 1-2 张命令行或容器状态截图，以增强运维说明的执行感。",
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            {"type": "bullet", "text": "查看后端日志：`docker compose logs backend --tail=50`。"},
            {"type": "bullet", "text": "查看异步 worker 日志：`docker compose logs search-indexer --tail=50`、`docker compose logs commerce-sync --tail=50`。"},
            {"type": "bullet", "text": "查看整体状态：`docker compose ps` 或 `make status`。"},
            {"type": "bullet", "text": "补写历史搜索事件：`make backfill-search-outbox`。"},
            {"type": "bullet", "text": "查看健康检查：`make health`。"},
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            {"type": "bullet", "text": "数据库连接失败：检查 mysql 容器是否 running、端口是否冲突、backend 是否等待依赖完成。"},
            {"type": "bullet", "text": "认证服务异常：重启 `supertokens` 与 `supertokens-db`，必要时执行 `make restart-auth`。"},
            {"type": "bullet", "text": "前端接口地址错误：检查根目录 `.env` 中 `VITE_*` 变量并重建前端镜像。"},
            {"type": "bullet", "text": "搜索结果未更新：检查 `search-indexer` 日志、OutboxEvent 状态，以及是否需要执行 `make backfill-search-outbox`。"},
            {"type": "bullet", "text": "Vendure 同步异常：检查 `commerce-sync` 与 `vendure` 日志，确认同步状态字段和错误信息。"},
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            {"type": "bullet", "text": "快速清空与重建：`docker compose down -v` 后重新执行 `make init`、`make start`。"},
            {"type": "bullet", "text": "业务数据重建：`npm run db:reset-and-seed-demo --prefix backend`。"},
            {"type": "bullet", "text": "搜索事件回填：在历史商品未补写事件时执行 `npm run db:backfill-search-outbox --prefix backend`。"},
            {"type": "bullet", "text": "当前未做独立自动化备份，课程演示阶段主要依赖脚本重建。"},
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            {"type": "bullet", "text": "重点巡检容器健康状态、后端错误日志、前端可访问性和数据库初始化结果。"},
            {"type": "bullet", "text": "消息实时性可通过消息页双窗口演示做快速人工巡检。"},
            {"type": "bullet", "text": "搜索、对象存储、认证和 Vendure 同步如异常，优先从对应容器日志定位。"},
            {"type": "bullet", "text": "对异步链路还应关注 Outbox 积压、失败状态和同步错误字段。"},
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            "日常运维以本地复现和演示稳定性为主，复杂的生产级监控、告警平台和自动备份体系不在本轮交付范围内。当前异步 worker 已独立运行，但监控和人工补偿能力仍偏轻量。",
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            {"type": "bullet", "text": "灰度发布在当前课程环境下等价于：先在本机或备用环境完成构建与启动，再切换演示入口。"},
            {"type": "bullet", "text": "建议先验证 `npm test --prefix backend` 与 `npm run build --prefix frontend`，再更新运行容器。"},
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            {"type": "bullet", "text": "若新版本异常，可回退到上一版仓库提交并重新 `make start`。"},
            {"type": "bullet", "text": "若数据结构已变化，应先确认是否需要执行 `make init-reset` 或演示数据重建。"},
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            {"type": "bullet", "text": "数据库结构调整以 Prisma schema 为主，通过 `db push` 或初始化脚本生效。"},
            {"type": "bullet", "text": "当前不维护独立 SQL migration 文档，因此所有结构变更必须同步更新 `schema.prisma` 和数据库设计说明书。"},
        ],
    )
    fill_text_cell(
        doc.tables[20].cell(0, 0),
        [
            "当前的版本升级与回滚策略更偏课程项目实践，重点是简单、可解释、可在答辩现场快速恢复运行状态。",
        ],
    )
    fill_text_cell(
        doc.tables[21].cell(0, 0),
        [
            {"type": "bullet", "text": "故障 1：后端接口不可达。现象：健康检查失败。排查：看 backend 与 mysql 日志。处置：重启后端或等待数据库就绪。预防：先执行 `make health`。"},
            {"type": "bullet", "text": "故障 2：前端打开但接口 404/跨域异常。排查：检查 `.env` 中 API 地址。处置：重建前端镜像。预防：改动变量后同步更新文档。"},
            {"type": "bullet", "text": "故障 3：图片上传失败。排查：检查 MinIO 容器与 bucket 初始化。处置：重启 MinIO 并查看媒体模块日志。预防：启动后先做一次发布页上传验证。"},
            {"type": "bullet", "text": "故障 4：消息实时推送不工作。排查：看 Socket.IO 连接与会话权限。处置：刷新页面或重启后端。预防：发布前做双窗口测试。"},
            {"type": "bullet", "text": "故障 5：搜索结果未及时更新。排查：检查 `search-indexer`、Meilisearch 与 OutboxEvent。处置：恢复 worker、必要时执行搜索事件回填。预防：启动后先验证一次发布与下单对搜索的影响。"},
            {"type": "bullet", "text": "故障 6：Vendure 同步失败。排查：检查 `commerce-sync`、`vendure`、同步状态字段和错误日志。处置：恢复容器后等待重试，必要时人工补偿。预防：演示前先做一次商品发布和订单创建联调。"},
            {"type": "bullet", "text": "故障 7：演示数据混乱。排查：确认是否误操作了生产态数据。处置：执行 `db:reset-and-seed-demo`。预防：区分空库基线与演示种子基线。"},
        ],
    )
    fill_text_cell(
        doc.tables[22].cell(0, 0),
        [
            {"type": "bullet", "text": "当前瓶颈主要在单体后端、数据库列表查询、消息实时性以及异步链路的观测能力。"},
            {"type": "bullet", "text": "后续扩展方向包括：补充更强缓存与搜索策略、完善 Outbox 监控补偿、拆分更细的治理/消息能力、增加前端自动化测试和正式监控。"},
            {"type": "bullet", "text": "课程阶段优先保证稳定复现，而非高并发扩容。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-08_部署与运维手册.docx")


def build_d09() -> None:
    doc = Document(TEMPLATE_DIR / "D-09_课程设计总结报告_模板.docx")
    fill_cover(doc.tables[0], code="D-09", name="课程设计总结报告")

    fill_text_cell(
        doc.tables[2].cell(0, 0),
        [
            "封面信息已按模板保留，目录建议在 Word 中打开文档后使用“引用 -> 目录 -> 自动目录”进行最终更新。本次生成版本已按正式章节结构组织内容，便于后续直接更新目录字段。",
        ],
    )
    fill_text_cell(
        doc.tables[3].cell(0, 0),
        [
            "中文摘要：SwapCampus 是一个面向校内师生的校园闲置物品交易平台，围绕商品浏览、发布、收藏、下单、评价、申诉、消息会话、校园服务和后台治理构建可运行的课程设计交付版本。项目采用 React、NestJS、MySQL、MinIO、Meilisearch、SuperTokens、Vendure 与 Docker Compose 技术路线，重点解决微信群交易信息分散、交易留痕不足、信用约束弱和部署复现困难等问题。当前系统不仅具备可部署、可演示、可测试和可追溯的主链路能力，也已经把搜索索引和外围电商同步改造成基于 Outbox 的独立异步处理链路。",
        ],
    )
    fill_text_cell(
        doc.tables[4].cell(0, 0),
        [
            "Abstract: SwapCampus is a campus-oriented idle goods trading platform for students and staff. The project delivers a runnable course-design version covering product browsing, publishing, favorites, ordering, reviews, appeals, messaging, campus services, and back-office governance. The system is built with React, NestJS, MySQL, MinIO, Meilisearch, SuperTokens, Vendure, and Docker Compose. It focuses on replacing fragmented group-chat trading with a traceable, trustworthy, and reproducible engineering solution, while decoupling search indexing and external commerce synchronization through an Outbox-based asynchronous workflow.",
        ],
    )
    fill_text_cell(
        doc.tables[5].cell(0, 0),
        [
            "本报告是在前序 D-01 到 D-08 文档基础上的最终汇总，用于从需求、设计、实现、测试、部署和项目管理多个视角说明 SwapCampus 的课程设计完成情况。",
        ],
    )
    fill_text_cell(
        doc.tables[6].cell(0, 0),
        [
            "校园闲置物品交易场景具有真实、频繁、低门槛的特点，但长期依赖微信群和熟人转发，导致信息沉底快、规则不可控、交易风险难追踪。SwapCampus 试图在明确的校内场景中，将身份、商品、交易、消息、信用和治理能力纳入同一系统之中，以展示软件工程在真实校园场景中的应用价值。",
        ],
    )
    fill_text_cell(
        doc.tables[7].cell(0, 0),
        [
            {"type": "bullet", "text": "完成商品发布、浏览、收藏、下单、评价和申诉闭环。"},
            {"type": "bullet", "text": "建立基于学号实名和信用分的可信体系。"},
            {"type": "bullet", "text": "提供消息会话、校园服务与后台治理能力。"},
            {"type": "bullet", "text": "形成可复现部署、可追溯文档和可演示验收材料。"},
        ],
    )
    fill_text_cell(
        doc.tables[8].cell(0, 0),
        [
            "全文结构分为需求分析、系统设计、系统实现、系统测试、系统部署与演示、项目管理总结以及结论与展望七个部分，既覆盖软件工程过程，也覆盖当前项目的真实运行状态。",
        ],
    )
    fill_text_cell(
        doc.tables[9].cell(0, 0),
        [
            {"type": "bullet", "text": "主要用户包括游客、普通用户和管理员。"},
            {"type": "bullet", "text": "典型场景包括：浏览闲置商品、校内面交、会话沟通、校园服务委托、后台审核和申诉处理。"},
            {"type": "bullet", "text": "系统重点服务教材、宿舍用品、数码配件、运动器材等校园高频闲置品类。"},
        ],
    )
    fill_text_cell(
        doc.tables[10].cell(0, 0),
        [
            {"type": "bullet", "text": "功能需求覆盖用户体系、商品能力、订单能力、消息能力、校园服务能力、信用与成长能力、后台治理能力。"},
            {"type": "bullet", "text": "当前主链路均已有实现与验收证据，可在追踪矩阵与测试报告中对应定位。"},
        ],
    )
    fill_text_cell(
        doc.tables[11].cell(0, 0),
        [
            {"type": "bullet", "text": "部署可复现：`make init` 与 `make start` 可完成完整环境启动。"},
            {"type": "bullet", "text": "数据口径清晰：默认空库建表，可选演示种子单独重建。"},
            {"type": "bullet", "text": "可维护性较好：文档与代码同仓，结构与实现同步迭代。"},
            {"type": "bullet", "text": "事件驱动副作用已落地：搜索索引和 Vendure 同步由独立 worker 异步收敛。"},
            {"type": "bullet", "text": "边界明确：完整 IM、正式压测和安全扫描仍为后续增强项。"},
        ],
    )
    fill_text_cell(
        doc.tables[12].cell(0, 0),
        [
            "系统总体架构采用前后端分离 + 主业务模块化单体 + 独立异步 worker + Compose 编排的方式，在课程规模下兼顾表达清晰和实现效率。浏览器访问 React 前端，前端通过 REST 与 Socket.IO 调用 NestJS 后端，后端在主事务中写入 OutboxEvent，再由 `search-indexer` 与 `commerce-sync` 分别把搜索索引和 Vendure 状态异步追平。",
        ],
    )
    fill_text_cell(
        doc.tables[13].cell(0, 0),
        [
            {"type": "bullet", "text": "关键模块包括：认证与用户、商品与收藏、订单与评价、消息、校园服务、信用中心、后台治理和媒体模块。"},
            {"type": "bullet", "text": "商品与订单模块承担核心交易闭环；消息与校园服务模块扩展真实校园使用场景；后台治理模块保障平台规则执行；Outbox、search-indexer 与 commerce-sync 负责异步副作用处理。"},
        ],
    )
    fill_text_cell(
        doc.tables[14].cell(0, 0),
        [
            {"type": "bullet", "text": "数据库采用 User、Product、Order、Conversation、CampusServiceListing 等核心表作为主干。"},
            {"type": "bullet", "text": "OutboxEvent、同步状态字段和外部系统 ID 用于支撑搜索索引与 Vendure 的最终一致。"},
            {"type": "bullet", "text": "订单快照、行为记录与信用流水帮助系统兼顾留痕、推荐和积分能力。"},
            {"type": "bullet", "text": "正式 ER 图仍建议作为附录补充。"},
        ],
    )
    fill_text_cell(
        doc.tables[15].cell(0, 0),
        [
            {"type": "bullet", "text": "关键接口覆盖认证、商品、订单、消息、校园服务、信用中心与后台治理。"},
            {"type": "bullet", "text": "接口风格以 REST 为主，消息实时追加通过 Socket.IO 完成。"},
        ],
    )
    fill_text_cell(
        doc.tables[16].cell(0, 0),
        [
            {"type": "bullet", "text": "算法层面重点不在复杂模型，而在状态流转控制、规则审核、推荐排序信号与幂等式写入。"},
            {"type": "bullet", "text": "这类算法更贴近课程项目对业务规则与工程实现结合的要求。"},
        ],
    )
    fill_text_cell(
        doc.tables[17].cell(0, 0),
        [
            "系统设计已能说明当前项目为何采用“模块化单体 + MySQL 主库 + SuperTokens 会话 + Compose 部署”的架构组合；正式答辩建议把架构图、ER 图、时序图作为附录一并补齐。",
        ],
    )
    fill_text_cell(
        doc.tables[18].cell(0, 0),
        [
            {"type": "bullet", "text": "开发环境：React 18、Vite、NestJS、Prisma、MySQL、MinIO、Meilisearch、SuperTokens、Docker Compose。"},
            {"type": "bullet", "text": "工具链：Git、npm、Makefile、GitHub Actions。"},
            {"type": "bullet", "text": "工程资产：README、课程文档、截图证据、验收记录与交付清单。"},
        ],
    )
    fill_text_cell(
        doc.tables[19].cell(0, 0),
        [
            {"type": "bullet", "text": "商品发布：多图上传、对象存储写入、本地规则 + LLM 审核后入库。"},
            {"type": "bullet", "text": "订单闭环：创建订单时保留商品快照，并自动建立消息上下文。"},
            {"type": "bullet", "text": "异步同步：商品、订单和卖家状态变化通过 OutboxEvent 分发给 `search-indexer` 与 `commerce-sync`。"},
            {"type": "bullet", "text": "消息能力：REST 读写 + Socket.IO 房间推送，满足演示级实时追加。"},
            {"type": "bullet", "text": "校园服务：采用 REQUEST/OFFER 双向意图和 listing + order 双层模型。"},
            {"type": "bullet", "text": "治理能力：后台支持商品审核、举报处理、申诉处理、实名审核与封禁。"},
        ],
    )
    fill_text_cell(
        doc.tables[20].cell(0, 0),
        [
            {"type": "bullet", "text": "关键决策 1：主业务保持 NestJS 模块化单体，同时把搜索索引和 Vendure 同步拆成独立异步 worker。"},
            {"type": "bullet", "text": "关键决策 2：采用 MySQL 作为业务主库，SuperTokens 独立维护认证库。"},
            {"type": "bullet", "text": "关键决策 3：先完成演示级实时消息、Outbox 解耦与部署复现，再逐步增强 IM、监控和自动化测试。"},
        ],
    )
    fill_text_cell(
        doc.tables[21].cell(0, 0),
        [
            "系统实现遵循“先打通链路、再补治理与证据、最后统一文档”的节奏，重点不是单点炫技，而是把课程要求对应到真实可运行的工程资产上。",
        ],
    )
    fill_text_cell(
        doc.tables[22].cell(0, 0),
        [
            {"type": "bullet", "text": "测试策略以自动化基础校验 + 人工主链路联调 + Docker 部署验证组成。"},
            {"type": "bullet", "text": "已执行的关键用例包括注册登录、商品发布、收藏、订单、消息、校园服务、后台治理、前端构建和后端测试。"},
        ],
    )
    fill_text_cell(
        doc.tables[23].cell(0, 0),
        [
            {"type": "bullet", "text": "后端测试命令、前端构建命令与 Compose 启动命令均已通过验证。"},
            {"type": "bullet", "text": "首页、登录、详情、发布、消息、个人中心和后台截图已归档。"},
            {"type": "bullet", "text": "系统已达到课程设计“可交付、可演示、可追溯”的要求。"},
        ],
    )
    fill_text_cell(
        doc.tables[24].cell(0, 0),
        [
            {"type": "bullet", "text": "缺陷主要集中在早期容器依赖注入、端口冲突、消息归属渲染和文档口径不一致等问题，均已修复或纳入边界说明。"},
            {"type": "bullet", "text": "当前质量风险主要来自 E2E、性能和安全专项尚未形成正式报告。"},
        ],
    )
    fill_text_cell(
        doc.tables[25].cell(0, 0),
        [
            "测试章节说明当前项目并非“没有问题”，而是已经把“已验证能力”和“未补齐专项”明确拆开，避免答辩时出现过度承诺。",
        ],
    )
    fill_text_cell(
        doc.tables[26].cell(0, 0),
        [
            "当前系统可通过 Docker Compose 独立部署并完成主要业务演示。以下截图为运行态证据：",
            {"type": "image", "path": IMAGE_DIR / "home-final.png", "caption": "图 5 首页运行态", "width": Cm(14.0)},
            {"type": "image", "path": IMAGE_DIR / "campus-services-final.png", "caption": "图 6 校园服务运行态", "width": Cm(14.0)},
            {"type": "image", "path": IMAGE_DIR / "messages-final.png", "caption": "图 7 消息页运行态", "width": Cm(14.0)},
            {"type": "image", "path": IMAGE_DIR / "admin-final.png", "caption": "图 8 后台运行态", "width": Cm(14.0)},
            "演示视频链接当前未嵌入，建议在 D-11 生成后补充视频地址或二维码。",
        ],
    )
    fill_text_cell(
        doc.tables[27].cell(0, 0),
        [
            {"type": "bullet", "text": "团队结构建议包括组长、前端负责人、后端负责人、测试文档负责人。"},
            {"type": "bullet", "text": "当前项目实际贡献集中在产品与视觉、后端与数据库、测试与运维三个维度。"},
        ],
    )
    fill_text_cell(
        doc.tables[28].cell(0, 0),
        [
            {"type": "bullet", "text": "阶段一：基础工程完成仓库结构、Compose、Prisma schema 与页面骨架。"},
            {"type": "bullet", "text": "阶段二：主链路完成商品流、注册登录、发布、订单与消息联动。"},
            {"type": "bullet", "text": "阶段三：治理与扩展完成后台审核、举报、封禁、校园服务、信用中心。"},
            {"type": "bullet", "text": "阶段四：对象存储、截图证据、文档归档与部署说明收口。"},
        ],
    )
    fill_text_cell(
        doc.tables[29].cell(0, 0),
        [
            {"type": "bullet", "text": "主要风险包括环境复杂度、功能范围膨胀、实时消息复杂度和答辩材料滞后。"},
            {"type": "bullet", "text": "应对策略包括范围控制、统一脚本入口、先交付可运行版本、把图示与增强项单独列清单补齐。"},
        ],
    )
    fill_text_cell(
        doc.tables[30].cell(0, 0),
        [
            {"type": "bullet", "text": "项目实践表明，边实现边验收边同步文档比最后集中补文档更可靠。"},
            {"type": "bullet", "text": "统一的事实源（源码、README、测试与验收记录）对于控制文档失真非常关键。"},
            {"type": "bullet", "text": "课程项目也可以用真实工程方式组织，而不是只做页面原型。"},
        ],
    )
    fill_text_cell(
        doc.tables[31].cell(0, 0),
        [
            "项目管理总结显示，本项目最大的收获之一是不再把文档当作附属品，而是作为设计、实现、测试和答辩之间的连接层。",
        ],
    )
    fill_text_cell(
        doc.tables[32].cell(0, 0),
        [
            {"type": "bullet", "text": "OpenAI Codex / GPT 系列：用于文档结构重组、模板填充、生成脚本编写与排版整理。"},
            {"type": "bullet", "text": "使用比例：约 20% 作为写作与整理辅助，不替代源码、测试和验收事实。"},
            {"type": "bullet", "text": "修订量：关键章节已结合当前源码、README 与文档记录做人工复核和改写。"},
            {"type": "bullet", "text": "负责人：待补充。"},
        ],
    )
    fill_text_cell(
        doc.tables[33].cell(0, 0),
        [
            "本项目已经完成了一个课程设计应有的主要工作：明确场景、建立需求、完成架构设计、落地数据库与代码实现、进行部署与验收，并形成成体系的文档交付。",
        ],
    )
    fill_text_cell(
        doc.tables[34].cell(0, 0),
        [
            {"type": "bullet", "text": "从软件工程角度看，项目已经体现了需求分析、模块划分、数据建模、测试验证、部署复现和文档治理等核心能力。"},
            {"type": "bullet", "text": "从产品角度看，项目已经不是静态页面集合，而是具备真实交互与治理链路的系统。"},
        ],
    )
    fill_text_cell(
        doc.tables[35].cell(0, 0),
        [
            {"type": "bullet", "text": "完整 IM 能力、正式图示材料、前端 E2E、性能与安全专项仍需增强。"},
            {"type": "bullet", "text": "若继续迭代，可进一步完善推荐排序、风控策略、运营看板与跨设备体验。"},
        ],
    )
    fill_text_cell(
        doc.tables[36].cell(0, 0),
        [
            {"type": "bullet", "text": "本项目强化了“需求、实现、验证、文档”四位一体的工程意识。"},
            {"type": "bullet", "text": "建议课程后续继续强调图示表达、自动化验证和文档追溯矩阵，以帮助学生形成更完整的工程闭环。"},
        ],
    )
    fill_text_cell(
        doc.tables[37].cell(0, 0),
        [
            "结论与展望部分说明：SwapCampus 已达到课程设计交付的核心目标，剩余主要工作并非“补出一个新系统”，而是把已有系统的图示、证据和正式材料进一步抛光。",
        ],
    )
    fill_text_cell(
        doc.tables[38].cell(0, 0),
        [
            {"type": "bullet", "text": "[1] React Documentation."},
            {"type": "bullet", "text": "[2] Vite Documentation."},
            {"type": "bullet", "text": "[3] NestJS Documentation."},
            {"type": "bullet", "text": "[4] Prisma Documentation."},
            {"type": "bullet", "text": "[5] MySQL 8 Reference Manual."},
            {"type": "bullet", "text": "[6] Docker Compose Documentation."},
            {"type": "bullet", "text": "[7] SuperTokens Documentation."},
            {"type": "bullet", "text": "[8] Meilisearch Documentation."},
            {"type": "bullet", "text": "[9] MinIO Documentation."},
            {"type": "bullet", "text": "[10] Socket.IO Documentation."},
            {"type": "bullet", "text": "[11] OWASP Web Security Testing Guide."},
        ],
    )
    fill_text_cell(
        doc.tables[39].cell(0, 0),
        [
            {"type": "bullet", "text": "附录建议补充：全局 ER 图、架构图、关键时序图、覆盖率截图、压测/安全扫描报告、完整 API 摘要。"},
            {"type": "bullet", "text": "当前可直接引用的证据包括：`artifacts/screenshots/` 截图、`docs/14-验收与演示记录.md`、`docs/17-任务需求实现追踪矩阵.md`。"},
        ],
    )

    doc.save(OUTPUT_DIR / "D-09_课程设计总结报告.docx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_d01()
    build_d02()
    build_d03()
    build_d04()
    build_d05()
    build_d06()
    build_d07()
    build_d08()
    build_d09()

    write_readme(
        [
            "D-01 第 6.1 节：建议补 1 张甘特图或项目计划截图。",
            "D-02 第 3.1 节：建议补 1 张总用例图与 1-2 张子用例图。",
            "D-03 第 2.2、3.2、4.1、4.2、5、6 节：建议补架构图、模块依赖图、时序图、状态机图、部署图、ER 图。",
            "D-04 第 5.1、5.2 节：建议补核心类图与关键时序图。",
            "D-05 第 2 节：建议补全局 ER 图与交易/消息/校园服务子图。",
            "D-06 第 5、6、7、8 节：建议补覆盖率截图、性能曲线、安全扫描输出和缺陷统计图。",
            "D-08 第 1 节：建议补正式部署拓扑图；第 3 节建议补命令执行截图。",
            "D-09 第 8、13 节：建议补部署架构图、演示视频链接以及附录图示材料。",
        ]
    )


if __name__ == "__main__":
    main()
